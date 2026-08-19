from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(r"D:\Projects\hongmen")
REFERENCE = ROOT / ".docx-work" / "reference.docx"
OUTPUT = ROOT / "忆程ProjectMemo_初赛作品说明文档_迁移稿.docx"
ASSETS = ROOT / "docs" / "competition" / "assets"

FONT = "宋体"
FOOTER_TEXT = "中国高校计算机大赛—人工智能创意赛（鸿蒙赛道）组委会编制"


def set_run_font(run, size=10.5, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def format_paragraph(p, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=0, first_line=True, keep=False):
    p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.keep_with_next = keep
    if first_line:
        pf.first_line_indent = Pt(21)
    else:
        pf.first_line_indent = Pt(0)
    return p


def add_text(doc, text, *, size=10.5, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             before=0, after=0, first_line=True, keep=False):
    p = doc.add_paragraph()
    format_paragraph(p, align=align, before=before, after=after, first_line=first_line, keep=keep)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12, 4: 10.5}
    before = {1: 8, 2: 6, 3: 4, 4: 3}[level]
    after = {1: 5, 2: 4, 3: 3, 4: 2}[level]
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.LEFT, before=before, after=after, first_line=False, keep=True)
    r = p.add_run(text)
    set_run_font(r, size=sizes[level], bold=True)
    return p


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    return p


def add_figure(doc, filename, caption, description=None, width_cm=15.0):
    path = ASSETS / filename
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.CENTER, after=3, first_line=False, keep=True)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    format_paragraph(cap, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=4, first_line=False, keep=bool(description))
    cr = cap.add_run(caption)
    set_run_font(cr, size=10.5, bold=False)
    if description:
        add_text(doc, description, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0, first_line=True)


def remove_instruction_tail(doc):
    marker = None
    for p in doc.paragraphs:
        if p.text.strip() == "作品说明文档提交规范说明":
            marker = p._element
            break
    if marker is None:
        raise RuntimeError("Cannot find submission-instruction marker")
    body = doc._element.body
    children = list(body)
    start = children.index(marker)
    for child in children[start:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_page_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, size=8)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    cached = OxmlElement("w:t")
    cached.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, cached, fld_end])


def clear_footer(footer):
    for p in list(footer.paragraphs):
        for child in list(p._element):
            p._element.remove(child)
    for t in list(footer.tables):
        t._element.getparent().remove(t._element)


def style_footer_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.tab_stops.add_tab_stop(Cm(15), WD_TAB_ALIGNMENT.RIGHT)


def add_footer_text(p, text):
    r = p.add_run(text)
    set_run_font(r, size=8)


def configure_page_system(doc):
    settings = doc.settings._element
    if settings.find(qn("w:evenAndOddHeaders")) is None:
        settings.append(OxmlElement("w:evenAndOddHeaders"))
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")

    for section in doc.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(3)
        section.header_distance = Cm(1.5)
        section.footer_distance = Cm(1.5)

        odd = section.footer
        clear_footer(odd)
        op = odd.paragraphs[0]
        style_footer_paragraph(op)
        add_footer_text(op, FOOTER_TEXT)
        add_footer_text(op, "\t")
        add_page_field(op)

        even = section.even_page_footer
        clear_footer(even)
        ep = even.paragraphs[0]
        style_footer_paragraph(ep)
        add_page_field(ep)
        add_footer_text(ep, "\t")
        add_footer_text(ep, FOOTER_TEXT)


def build():
    doc = Document(REFERENCE)
    remove_instruction_tail(doc)
    configure_page_system(doc)
    doc.core_properties.title = "忆程 ProjectMemo 初赛作品说明文档"
    doc.core_properties.subject = "2026中国高校计算机大赛—人工智能创意赛（鸿蒙赛道）Agent创新"
    doc.core_properties.author = "万卷siu"

    add_text(doc, "忆程 ProjectMemo", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             before=0, after=5, first_line=False, keep=True)
    add_text(doc, "万卷siu", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
             before=0, after=12, first_line=False, keep=True)

    add_heading(doc, "一、创意描述", 1)
    add_text(doc, "将项目碎片转化为会主动推动交付的长期记忆", size=10.5,
             align=WD_ALIGN_PARAGRAPH.CENTER, before=2, after=10, first_line=False)

    add_heading(doc, "二、设计稿（Agent创新）", 1)
    add_heading(doc, "（一）设计概述", 2)
    add_text(doc,
             "忆程 ProjectMemo 面向大学生课程设计、科研训练和学科竞赛。系统将论文笔记、实验结果、代码问题、会议结论与材料要求沉淀为可关联、可追溯的项目记忆，并在截止逼近、风险未处理、项目停滞或材料缺失时主动介入。用户确认后，建议转化为行动；行动完成结果再回写为复盘记忆，形成持续推进闭环。",
             after=5)
    add_figure(doc, "图1_ProjectMemo_目标用户流程.drawio.png",
               "图1  忆程 ProjectMemo 目标用户流程",
               "用户从项目碎片输入开始，经过结构化、历史关联和风险识别，最终完成用户确认、行动回写与成果生成。")

    add_page_break(doc)
    add_heading(doc, "（二）Agent Workflow", 2)
    add_figure(doc, "图2_ProjectMemo_Agent_Workflow.drawio.png",
               "图2  忆程 ProjectMemo Agent Workflow",
               "系统按碎片输入、情境变化和成果请求进行路由。模型只承担结构化与内容生成节点；输入校验、规则判断、数据写入和事务控制由应用层负责，异常时回退确定性 Mock。")

    add_page_break(doc)
    add_heading(doc, "（三）主动介入与行动闭环", 2)
    add_figure(doc, "图3_ProjectMemo_主动介入与行动闭环.drawio.png",
               "图3  主动介入—行动—反馈—记忆闭环",
               "Agent 先识别触发情境，再组织引用证据和建议行动；接受、稍后或忽略均由用户决定。当前单用户 Demo 已实现行动标题、说明、优先级、截止日期与完成回执；图中“负责人”为后续团队协作扩展位。")

    add_page_break(doc)
    add_heading(doc, "（四）当前可运行界面", 2)
    add_figure(doc, "首页.png", "图4  产品价值主张与闭环入口（本地 Mock Demo）",
               "首页在首屏说明“碎片记录—项目记忆—主动介入—行动闭环—成果生成”的完整价值链，并明确标注离线 Mock 演示模式。")

    add_page_break(doc)
    add_figure(doc, "项目页面.png", "图5  项目空间与参赛项目状态（本地 Mock Demo）",
               "项目空间提供搜索、场景筛选和“需要关注”排序，并在卡片中集中展示提醒、待办、知识卡片、成果和参赛准备度。")

    add_page_break(doc)
    add_figure(doc, "项目详细.png", "图6  项目工作台与记录入口（本地 Mock Demo）",
               "项目工作台把截止状态、当前提醒、未完成行动和准备度置于顶部，通过页面大纲快速跳转到记录、提醒、行动、知识资产和成果区域。")

    add_page_break(doc)
    add_figure(doc, "主动介入.png", "图7  主动介入、证据说明与确认行动（本地 Mock Demo）",
               "主动提醒展示触发情境和建议行动，用户可接受、稍后或忽略。接受后创建行动项；行动完成结果自动沉淀为阶段复盘卡片。")

    add_page_break(doc)
    add_figure(doc, "成果呈现.png", "图8  从项目记忆生成并编辑成果（本地 Mock Demo）",
               "成果文档室可从已有知识卡片生成六类交付材料，并支持 Markdown 渲染预览、源码编辑、复制、导出与历史版本留存。")

    add_page_break(doc)
    add_heading(doc, "三、作品介绍", 1)
    intro = [
        "大学生在课程设计、科研训练和学科竞赛中，会持续产生论文笔记、实验结果、代码问题、会议结论和材料要求。这些信息散落在聊天、文档和个人笔记里，项目越久，越难回答“为何这样决定、风险是否处理、下一步做什么”。临近截止时，团队往往重新翻找记录、整理进展、拼装材料。通用问答助手擅长一次性生成，却缺少稳定的项目记忆、可追溯依据和持续推进机制。",
        "忆程 ProjectMemo 将 AI 从“被动回答”转变为“主动推动交付”的项目记忆 Agent。用户输入自然语言碎片后，系统提取摘要、关键词、任务、重要性和下一步，关联历史卡片并保留原文；再根据截止逼近、风险未处理、项目停滞、实验或材料缺口主动介入。提醒展示触发原因与引用证据，用户可接受、稍后或忽略；接受后生成行动项，完成结果自动沉淀为复盘卡片，形成“记录—关联—判断—行动—反馈—记忆”闭环。",
        "当前已完成可本地运行、完全离线演示的 Web Demo，支持碎片结构化、知识关联、行动板、主动提醒、带引用的记忆副驾驶和效果指标，并能生成周报、作品说明、答辩 PPT、README、简历描述和行动计划。成果支持 Markdown 渲染与编辑、复制、导出和版本留存。系统采用 Next.js、TypeScript、Prisma 与 SQLite；模型通过 OpenAI-compatible 接口接入，结构化结果需严格校验，超时、非法 JSON 或无密钥时自动回退确定性 Mock，保证演示稳定。",
        "忆程 ProjectMemo 的创新不在于增加一个笔记库，而在于让项目记忆成为可解释、可确认、可回写的主动服务。它帮助学生保留决策脉络，把过程材料转化为可复用知识资产和交付成果。后续可扩展语音、图片、后台提醒与团队协作，并将核心工具封装为小艺 Workflow；当前仅保留接口设计，不宣称已接入。",
    ]
    for paragraph in intro:
        add_text(doc, paragraph, size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, first_line=True)

    add_heading(doc, "说明", 4)
    add_text(doc,
             "图4至图8均来自本地可运行的离线 Mock 演示环境，使用模拟项目数据，不代表真实用户规模、官方审核结果或小艺已接入状态。",
             size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=0, first_line=True)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
