from pathlib import Path
from docx import Document

path = Path(r"D:\Projects\hongmen\.docx-work\reference.docx")
doc = Document(path)

print("PARAGRAPHS")
for i, p in enumerate(doc.paragraphs):
    text = p.text.replace("\t", "<TAB>").replace("\n", "<NL>")
    if text.strip():
        pf = p.paragraph_format
        print(f"P{i:03d} style={p.style.name!r} align={p.alignment} text={text!r}")

print("\nTABLES")
for ti, table in enumerate(doc.tables):
    print(f"TABLE {ti} rows={len(table.rows)} cols={len(table.columns)} style={table.style.name if table.style else None!r}")
    for ri, row in enumerate(table.rows):
        values = []
        for cell in row.cells:
            values.append(" | ".join(p.text for p in cell.paragraphs).replace("\n", "<NL>"))
        print(f"  R{ri:02d}: {values!r}")

print("\nSECTIONS")
for i, s in enumerate(doc.sections):
    print(i, s.page_width, s.page_height, s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)

print("\nSTYLES")
for s in doc.styles:
    if s.type == 1 and (s.name in {"Normal", "List Paragraph"} or "Title" in s.name or "Heading" in s.name):
        print(s.name, s.font.name, s.font.size, s.font.bold, s.font.color.rgb if s.font.color and s.font.color.type else None)
